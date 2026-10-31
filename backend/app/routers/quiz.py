"""
Quiz & Participant Analytics Router
NSS MJCET Nasha Mukti Abhiyaan
Server-side authoritative scoring, participant attempt recording, and Excel export.
"""

from typing import List, Dict, Any, Optional
from datetime import datetime, timezone
import io
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorDatabase
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.app.core.permissions import require_super_admin, require_editor_or_above
from backend.app.db.mongodb import get_database
from backend.app.schemas.quiz import (
    QuizConfigItem,
    QuizConfigUpdate,
    QuizQuestionPublicResponse,
    QuizQuestionAdminResponse,
    QuizQuestionCreate,
    QuizQuestionUpdate,
    QuizAttemptSubmit,
    QuizAttemptResult,
    QuestionBreakdownItem,
    LeaderboardEntryPublic,
    LeaderboardPublicResponse,
    LeaderboardSummaryItem,
    LeaderboardEntryAdmin,
    LeaderboardAdminResponse,
    LeaderboardSettingsUpdate,
    AttemptOverrideRequest,
)
from backend.app.utils.audit import record_audit_log

router = APIRouter(tags=["Quiz"])


def sanitize_excel_cell(val: Any) -> Any:
    """Neutralizes Excel formula injection risks for user-supplied strings."""
    if isinstance(val, str) and len(val) > 0:
        if val[0] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + val
    return val


# Public Endpoints
@router.get("/quiz/portal", response_model=List[QuizConfigItem])
@router.get("/quiz/list", response_model=List[QuizConfigItem])
async def list_public_quizzes(
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Returns all configured campaign quizzes for the Unified Quiz Portal.
    Includes live question count and points calculated from database.
    """
    cursor = db["quizzes"].find().sort("displayOrder", 1)
    quizzes = []
    async for q in cursor:
        q_id = q["id"]
        meta = resolve_quiz_meta(q_id)
        query_ids = meta.get("aliases", [q_id])

        q_count = await db["quiz_questions"].count_documents({
            "quizId": {"$in": query_ids},
            "publicationStatus": "published",
        })
        pipeline = [
            {"$match": {"quizId": {"$in": query_ids}, "publicationStatus": "published"}},
            {"$group": {"_id": None, "total": {"$sum": "$points"}}},
        ]
        pts_agg = await db["quiz_questions"].aggregate(pipeline).to_list(1)
        total_pts = pts_agg[0]["total"] if pts_agg else (q_count * 10)

        quizzes.append(
            QuizConfigItem(
                id=q_id,
                title=q.get("title", meta["title"]),
                slug=q.get("slug", meta["slug"]),
                description=q.get("description", ""),
                badge=q.get("badge", "Official NSS Challenge"),
                estimatedMinutes=q.get("estimatedMinutes", 10),
                status=q.get("status", "active"),
                theme=q.get("theme", "awareness"),
                passingScorePercentage=q.get("passingScorePercentage", 70),
                timeLimitMinutes=q.get("timeLimitMinutes", 15),
                questionCount=q_count,
                totalPoints=total_pts,
            )
        )
    return quizzes


@router.get("/quiz/public", response_model=List[QuizQuestionPublicResponse])
@router.get("/quiz/public/{quiz_id}", response_model=List[QuizQuestionPublicResponse])
async def get_public_quiz_questions(
    quiz_id: Optional[str] = None,
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Returns public questions for participant attempts for a specific quiz.
    STRICT SECURITY: Excludes 'correctAnswerIds' and 'explanation' to prevent cheating.
    """
    target_id = (quiz_id or "nmba-awareness-challenge-v1").strip().lower()
    meta = resolve_quiz_meta(target_id)
    query_ids = meta.get("aliases", [meta["id"]])

    cursor = db["quiz_questions"].find({
        "quizId": {"$in": query_ids},
        "publicationStatus": "published",
    }).sort("displayOrder", 1)

    questions = []
    async for doc in cursor:
        doc["_id"] = str(doc["_id"])
        # Ensure correct answers and explanations are stripped
        doc.pop("correctAnswerIds", None)
        doc.pop("explanation", None)
        questions.append(QuizQuestionPublicResponse(**doc))
    return questions


@router.post("/quiz/submit", response_model=QuizAttemptResult)
async def submit_quiz_attempt(
    payload: QuizAttemptSubmit,
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Authoritative server-side quiz evaluation and participant attempt recording.
    Calculates genuine score against canonical database records for the specific quiz
    and returns result breakdown.
    """
    target_id = payload.quizId.strip().lower()
    meta = resolve_quiz_meta(target_id)
    query_ids = meta.get("aliases", [meta["id"]])

    # Fetch authoritative published questions with answers for this quiz
    cursor = db["quiz_questions"].find({
        "quizId": {"$in": query_ids},
        "publicationStatus": "published",
    }).sort("displayOrder", 1)

    canonical_questions = {}
    async for q in cursor:
        canonical_questions[q["id"]] = q

    if not canonical_questions:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=f"Quiz questions for '{meta['title']}' are temporarily unavailable.",
        )

    earned_score = 0
    max_score = 0
    correct_count = 0
    incorrect_count = 0
    question_breakdown: List[QuestionBreakdownItem] = []

    for q_id, q in canonical_questions.items():
        points = q.get("points", 10)
        max_score += points
        correct_ids = q.get("correctAnswerIds", [])
        selected_ids = payload.userAnswers.get(q_id, [])

        is_correct = False
        if q.get("type") == "multiple":
            # Multi-select: exact match of set
            is_correct = set(selected_ids) == set(correct_ids) and len(selected_ids) > 0
        else:
            # Single or boolean
            is_correct = len(selected_ids) == 1 and len(correct_ids) >= 1 and selected_ids[0] == correct_ids[0]

        pts = points if is_correct else 0
        earned_score += pts

        if is_correct:
            correct_count += 1
        else:
            incorrect_count += 1

        question_breakdown.append(
            QuestionBreakdownItem(
                questionId=q_id,
                questionText=q.get("text", ""),
                category=q.get("category", ""),
                type=q.get("type", "single"),
                points=points,
                earnedPoints=pts,
                selectedOptionIds=selected_ids,
                correctOptionIds=correct_ids,
                isCorrect=is_correct,
                options=q.get("options", []),
                explanation=q.get("explanation", ""),
            )
        )

    percentage = round((earned_score / max_score) * 100) if max_score > 0 else 0
    passed = percentage >= 70

    # Assign quiz-specific badge and feedback message
    is_unity = meta["id"] == "national-unity-day-quiz-v1"
    if is_unity:
        if percentage >= 90:
            badge_name = "Rashtriya Ekta Champion"
            badge_desc = "Exemplary understanding of national integration, constitutional duties, and Sardar Patel's legacy."
            summary_msg = "Outstanding! You have demonstrated exceptional mastery of India's unity, constitutional values, and civic heritage."
        elif percentage >= 70:
            badge_name = "Patriotic Civic Ambassador"
            badge_desc = "Distinguished appreciation of Indian unity in diversity and selfless community service."
            summary_msg = "Well done! You embody the spirit of national unity and collegiate civic responsibility."
        else:
            badge_name = "National Unity Contributor"
            badge_desc = "Engaged participant exploring India's heritage, integration, and civic responsibilities."
            summary_msg = "Knowledge inspires nation building. Review the question explanations below to deepen your understanding of India's heritage."
    else:
        if percentage >= 90:
            badge_name = "Nasha Mukti Campus Champion"
            badge_desc = "Exceptional mastery of prevention science, refusal agency, and community care."
            summary_msg = "Outstanding! You have demonstrated exceptional awareness of substance dependency science and compassionate peer support."
        elif percentage >= 70:
            badge_name = "Informed Peer Advocate"
            badge_desc = "Solid grasp of neurobiology, de-escalation, and verified helplines."
            summary_msg = "Well done! You possess the essential knowledge to make conscious choices and support peers in need."
        else:
            badge_name = "Awareness Contributor"
            badge_desc = "Completed the curriculum review and engaged in conscious learning."
            summary_msg = "Knowledge grows with reflection. Review the question explanations below to deepen your awareness of recovery science."

    now_iso = datetime.now(timezone.utc).isoformat()
    attempt_id = f"atmpt-{int(datetime.now().timestamp() * 1000)}"

    attempt_doc = {
        "attemptId": attempt_id,
        "quizId": meta["id"],
        "participant": payload.participant.model_dump(),
        "score": earned_score,
        "maxScore": max_score,
        "percentage": percentage,
        "passed": passed,
        "totalQuestions": len(canonical_questions),
        "correctCount": correct_count,
        "incorrectCount": incorrect_count,
        "badgeName": badge_name,
        "badgeDescription": badge_desc,
        "summaryMessage": summary_msg,
        "submittedAt": now_iso,
        "userAnswers": payload.userAnswers,
        "timeTakenSeconds": payload.timeTakenSeconds,
    }

    # Save to MongoDB quiz_attempts collection
    await db["quiz_attempts"].insert_one(attempt_doc)

    return QuizAttemptResult(
        attemptId=attempt_id,
        quizId=meta["id"],
        participantName=payload.participant.fullName.strip(),
        participantInstitution=payload.participant.institutionName.strip(),
        submittedAt=now_iso,
        score=earned_score,
        maxScore=max_score,
        percentage=percentage,
        passed=passed,
        totalQuestions=len(canonical_questions),
        correctCount=correct_count,
        incorrectCount=incorrect_count,
        badgeName=badge_name,
        badgeDescription=badge_desc,
        summaryMessage=summary_msg,
        questionBreakdown=question_breakdown,
        timeTakenSeconds=payload.timeTakenSeconds,
    )


# Protected CMS Endpoints for Quiz Management
@router.get("/admin/quizzes", response_model=List[QuizConfigItem])
async def list_admin_quizzes(
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Lists all quizzes with configuration settings and live statistics for CMS."""
    cursor = db["quizzes"].find().sort("displayOrder", 1)
    quizzes = []
    async for q in cursor:
        q_id = q["id"]
        meta = resolve_quiz_meta(q_id)
        query_ids = meta.get("aliases", [q_id])

        q_count = await db["quiz_questions"].count_documents({"quizId": {"$in": query_ids}})
        pipeline = [
            {"$match": {"quizId": {"$in": query_ids}}},
            {"$group": {"_id": None, "total": {"$sum": "$points"}}},
        ]
        pts_agg = await db["quiz_questions"].aggregate(pipeline).to_list(1)
        total_pts = pts_agg[0]["total"] if pts_agg else (q_count * 10)

        quizzes.append(
            QuizConfigItem(
                id=q_id,
                title=q.get("title", meta["title"]),
                slug=q.get("slug", meta["slug"]),
                description=q.get("description", ""),
                badge=q.get("badge", "Official NSS Challenge"),
                estimatedMinutes=q.get("estimatedMinutes", 10),
                status=q.get("status", "active"),
                theme=q.get("theme", "awareness"),
                passingScorePercentage=q.get("passingScorePercentage", 70),
                timeLimitMinutes=q.get("timeLimitMinutes", 15),
                questionCount=q_count,
                totalPoints=total_pts,
            )
        )
    return quizzes


@router.put("/admin/quizzes/{quiz_id}", response_model=QuizConfigItem)
async def update_admin_quiz_config(
    quiz_id: str,
    payload: QuizConfigUpdate,
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Updates quiz configuration (title, description, badge, duration, status, pass percentage).
    STRICT RBAC: Editors and Super Admins.
    """
    meta = resolve_quiz_meta(quiz_id)
    now_iso = datetime.now(timezone.utc).isoformat()

    update_dict = {k: v for k, v in payload.model_dump().items() if v is not None}
    update_dict["updatedAt"] = now_iso

    await db["quizzes"].update_one(
        {"id": meta["id"]},
        {"$set": update_dict},
        upsert=True,
    )

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="update_quiz_config",
        module="quizzes",
        record_id=meta["id"],
        metadata={"changes": list(update_dict.keys())},
    )

    q = await db["quizzes"].find_one({"id": meta["id"]})
    query_ids = meta.get("aliases", [meta["id"]])
    q_count = await db["quiz_questions"].count_documents({"quizId": {"$in": query_ids}})
    pipeline = [
        {"$match": {"quizId": {"$in": query_ids}}},
        {"$group": {"_id": None, "total": {"$sum": "$points"}}},
    ]
    pts_agg = await db["quiz_questions"].aggregate(pipeline).to_list(1)
    total_pts = pts_agg[0]["total"] if pts_agg else (q_count * 10)

    return QuizConfigItem(
        id=meta["id"],
        title=q.get("title", meta["title"]),
        slug=q.get("slug", meta["slug"]),
        description=q.get("description", ""),
        badge=q.get("badge", "Official NSS Challenge"),
        estimatedMinutes=q.get("estimatedMinutes", 10),
        status=q.get("status", "active"),
        theme=q.get("theme", "awareness"),
        passingScorePercentage=q.get("passingScorePercentage", 70),
        timeLimitMinutes=q.get("timeLimitMinutes", 15),
        questionCount=q_count,
        totalPoints=total_pts,
    )


# Protected CMS Endpoints for Quiz Questions
@router.get("/admin/quiz/questions", response_model=List[QuizQuestionAdminResponse])
async def list_admin_quiz_questions(
    quiz_id: Optional[str] = Query(None),
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Retrieves questions including correct answers and explanations for CMS.
    Supports filtering by specific quiz (e.g. ?quiz_id=national-unity-day-quiz-v1).
    """
    filter_query: Dict[str, Any] = {}
    if quiz_id and quiz_id.strip() and quiz_id.strip() != "all":
        meta = resolve_quiz_meta(quiz_id)
        query_ids = meta.get("aliases", [meta["id"]])
        filter_query["quizId"] = {"$in": query_ids}

    cursor = db["quiz_questions"].find(filter_query).sort("displayOrder", 1)
    results = []
    async for doc in cursor:
        doc["_id"] = str(doc["_id"])
        if "quizId" not in doc or not doc["quizId"]:
            doc["quizId"] = "nmba-awareness-challenge-v1"
        results.append(QuizQuestionAdminResponse(**doc))
    return results


@router.post("/admin/quiz/questions", response_model=QuizQuestionAdminResponse, status_code=status.HTTP_201_CREATED)
async def create_quiz_question(
    payload: QuizQuestionCreate,
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Creates a new quiz question under the specified quiz."""
    now_iso = datetime.now(timezone.utc).isoformat()
    q_id = f"q-{int(datetime.now().timestamp() * 1000) % 1000000}"

    doc = payload.model_dump()
    doc["id"] = q_id
    if not doc.get("quizId"):
        doc["quizId"] = "nmba-awareness-challenge-v1"
    doc["createdAt"] = now_iso
    doc["updatedAt"] = now_iso

    await db["quiz_questions"].insert_one(doc)
    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="create_quiz_question",
        module="quiz_questions",
        record_id=q_id,
        metadata={"category": payload.category, "quizId": doc["quizId"]},
    )

    doc["_id"] = str(doc.get("_id", ""))
    return QuizQuestionAdminResponse(**doc)


@router.put("/admin/quiz/questions/{id}", response_model=QuizQuestionAdminResponse)
async def update_quiz_question(
    id: str,
    payload: QuizQuestionUpdate,
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Updates a quiz question, options, answers, or explanation."""
    target = await db["quiz_questions"].find_one({"id": id})
    if not target:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Question '{id}' not found.")

    update_dict = {k: v for k, v in payload.model_dump().items() if v is not None}
    update_dict["updatedAt"] = datetime.now(timezone.utc).isoformat()

    await db["quiz_questions"].update_one({"id": id}, {"$set": update_dict})
    updated_doc = await db["quiz_questions"].find_one({"id": id})

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="update_quiz_question",
        module="quiz_questions",
        record_id=id,
        metadata={"changes": list(update_dict.keys())},
    )

    updated_doc["_id"] = str(updated_doc["_id"])
    return QuizQuestionAdminResponse(**updated_doc)


@router.delete("/admin/quiz/questions/{id}", status_code=status.HTTP_200_OK)
async def delete_quiz_question(
    id: str,
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Deletes a quiz question."""
    target = await db["quiz_questions"].find_one({"id": id})
    if not target:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Question '{id}' not found.")

    await db["quiz_questions"].delete_one({"id": id})
    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="delete_quiz_question",
        module="quiz_questions",
        record_id=id,
        metadata={"text": target.get("text")[:40]},
    )

    return {"message": f"Question '{id}' has been deleted successfully."}


# Protected Quiz Attempts & Excel Export
@router.get("/admin/quiz/attempts")
async def list_quiz_attempts(
    limit: int = Query(50, ge=1, le=200),
    skip: int = Query(0, ge=0),
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Lists real participant attempts from MongoDB with pagination."""
    cursor = db["quiz_attempts"].find().sort("submittedAt", -1).skip(skip).limit(limit)
    total = await db["quiz_attempts"].count_documents({})
    items = []
    async for doc in cursor:
        doc["_id"] = str(doc["_id"])
        items.append(doc)
    return {"total": total, "items": items, "limit": limit, "skip": skip}


@router.get("/admin/quiz/export")
async def export_quiz_attempts_excel(
    current_admin: Dict[str, Any] = Depends(require_super_admin),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Generates and returns an authentic .xlsx Excel workbook containing all participant
    quiz attempts and aggregate statistics using OpenPyXL.
    """
    cursor = db["quiz_attempts"].find().sort("submittedAt", -1)
    attempts = []
    async for doc in cursor:
        attempts.append(doc)

    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")  # Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # 1. Sheet 1: Quiz Attempts
    ws1 = wb.active
    ws1.title = "Quiz Attempts"

    headers1 = [
        "Attempt ID",
        "Submitted At (UTC)",
        "Participant Name",
        "Email",
        "Institution / College",
        "Roll Number",
        "Mobile Number",
        "Department",
        "Year of Study",
        "Score",
        "Max Score",
        "Percentage (%)",
        "Status",
        "Correct Answers",
        "Incorrect Answers",
        "Assigned Badge",
    ]
    ws1.append(headers1)
    for col_num in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for att in attempts:
        p = att.get("participant", {})
        row = [
            sanitize_excel_cell(att.get("attemptId", "")),
            sanitize_excel_cell(att.get("submittedAt", "")),
            sanitize_excel_cell(p.get("fullName", "")),
            sanitize_excel_cell(p.get("email", "")),
            sanitize_excel_cell(p.get("institutionName", "")),
            sanitize_excel_cell(p.get("rollNumber", "")),
            sanitize_excel_cell(p.get("phone", "")),
            sanitize_excel_cell(p.get("department", "")),
            sanitize_excel_cell(p.get("yearOfStudy", "")),
            att.get("score", 0),
            att.get("maxScore", 0),
            att.get("percentage", 0),
            "PASSED" if att.get("passed") else "ATTEMPTED",
            att.get("correctCount", 0),
            att.get("incorrectCount", 0),
            sanitize_excel_cell(att.get("badgeName", "")),
        ]
        ws1.append(row)

    # Apply data borders and alignment
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=1, max_col=len(headers1)):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    # Auto-fit column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 2. Sheet 2: Participant Summary
    ws2 = wb.create_sheet(title="Campaign Metrics")
    ws2.append(["Metric", "Recorded Value"])
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=1).font = header_font
    ws2.cell(row=1, column=2).fill = header_fill
    ws2.cell(row=1, column=2).font = header_font

    total_attempts = len(attempts)
    unique_emails = len(set(a.get("participant", {}).get("email") for a in attempts if a.get("participant", {}).get("email")))
    passed_count = sum(1 for a in attempts if a.get("passed"))
    avg_score = round(sum(a.get("percentage", 0) for a in attempts) / total_attempts, 1) if total_attempts > 0 else 0
    pass_rate = round((passed_count / total_attempts) * 100, 1) if total_attempts > 0 else 0

    metrics_rows = [
        ("Total Quiz Attempts", total_attempts),
        ("Unique Participants", unique_emails),
        ("Successful Completions (>= 70%)", passed_count),
        ("Average Score Percentage", f"{avg_score}%"),
        ("Passing Rate", f"{pass_rate}%"),
        ("Export Generated At (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Export Requested By", current_admin.get("email")),
    ]
    for m, val in metrics_rows:
        ws2.append([m, val])

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 30

    # Save to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="export_quiz_attempts_excel",
        module="quiz_analytics",
        record_id="export_all",
        metadata={"totalRecordsExported": total_attempts},
    )

    filename = f"nmba_quiz_attempts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# =========================================================================
# Quiz Leaderboard Helpers & Endpoints
# =========================================================================

QUIZ_REGISTRY = {
    "nmba-awareness-challenge-v1": {
        "id": "nmba-awareness-challenge-v1",
        "title": "Nasha Mukti Awareness Quiz",
        "slug": "nasha-mukti-awareness",
        "aliases": ["nmba-awareness-challenge-v1", "nasha-mukti-awareness", "awareness"],
    },
    "national-unity-day-quiz-v1": {
        "id": "national-unity-day-quiz-v1",
        "title": "National Unity Day Quiz",
        "slug": "national-unity-day",
        "aliases": ["national-unity-day-quiz-v1", "national-unity-day", "unity", "unity-day"],
    },
    "unified-nss-mjcet-leaderboard": {
        "id": "unified-nss-mjcet-leaderboard",
        "title": "Unified NSS MJCET Leaderboard",
        "slug": "unified",
        "aliases": ["unified-nss-mjcet-leaderboard", "unified", "combined", "all"],
    },
}


def resolve_quiz_meta(identifier: str) -> Dict[str, Any]:
    """Resolves quiz identifier or slug to standard quiz metadata."""
    clean_id = identifier.strip().lower()
    for meta in QUIZ_REGISTRY.values():
        if clean_id == meta["id"].lower() or clean_id == meta["slug"].lower() or clean_id in [a.lower() for a in meta["aliases"]]:
            return meta
    return {
        "id": identifier,
        "title": identifier.replace("-", " ").title(),
        "slug": identifier.lower(),
        "aliases": [identifier],
    }


def format_duration_string(seconds: Optional[int]) -> str:
    """Formats seconds into readable MM:SS or human string."""
    if seconds is None or seconds <= 0:
        return "—"
    m = seconds // 60
    s = seconds % 60
    if m > 0:
        return f"{m}m {s:02d}s" if s > 0 else f"{m}m"
    return f"{s}s"


def mask_participant_name(name: str) -> str:
    """Masks participant name for privacy when configured."""
    parts = name.strip().split()
    if not parts:
        return "Participant"
    masked = []
    for p in parts:
        if len(p) <= 2:
            masked.append(p[0] + "*")
        else:
            masked.append(p[0] + ("*" * (len(p) - 1)))
    return " ".join(masked)


async def get_or_create_leaderboard_record(
    db: AsyncIOMotorDatabase,
    quiz_meta: Dict[str, Any],
) -> Dict[str, Any]:
    """Retrieves or initializes the persistent leaderboard configuration document."""
    quiz_id = quiz_meta["id"]
    record = await db["quiz_leaderboards"].find_one({"quizId": quiz_id})
    if not record:
        now_iso = datetime.now(timezone.utc).isoformat()
        record = {
            "quizId": quiz_id,
            "quizTitle": quiz_meta["title"],
            "quizSlug": quiz_meta["slug"],
            "isPublished": False,
            "status": "coming_soon",
            "publishedAt": None,
            "publishedBy": None,
            "announcementDate": None,
            "showPublicNames": True,
            "highlightTopThree": True,
            "eligibilityRule": "any",
            "customNote": "Official certified rankings authenticated by NSS MJCET Examination Authority.",
            "createdAt": now_iso,
            "updatedAt": now_iso,
        }
        await db["quiz_leaderboards"].insert_one(record)
    return record


async def compute_unified_rankings(
    db: AsyncIOMotorDatabase,
    cfg: Dict[str, Any],
) -> List[Dict[str, Any]]:
    """
    Computes authoritative combined rankings across both quizzes:
    1. Aggregates attempts per participant email.
    2. Takes the best score and lowest time for each quiz.
    3. Sums combinedScore and totalTimeSeconds.
    4. Applies eligibility rule ('any' or 'both').
    5. Standard Olympic ranking based on (-combinedScore, totalTimeSeconds, latest_submittedAt).
    """
    cursor = db["quiz_attempts"].find()
    attempts = []
    async for a in cursor:
        attempts.append(a)

    eligibility = cfg.get("eligibilityRule", "any")

    # Group attempts by normalized participant email (or fallback to rollNumber / name)
    participants_map: Dict[str, Dict[str, Any]] = {}

    nmba_aliases = [a.lower() for a in QUIZ_REGISTRY["nmba-awareness-challenge-v1"]["aliases"]]
    unity_aliases = [a.lower() for a in QUIZ_REGISTRY["national-unity-day-quiz-v1"]["aliases"]]

    for a in attempts:
        p = a.get("participant", {})
        email = (p.get("email") or "").strip().lower()
        roll = (p.get("rollNumber") or "").strip().upper()
        name = (p.get("fullName") or "Participant").strip()

        key = email if email else (roll if roll else name.lower())
        if not key:
            continue

        if key not in participants_map:
            participants_map[key] = {
                "participant": p,
                "nmba_best": None,
                "unity_best": None,
            }

        q_id = (a.get("quizId") or "nmba-awareness-challenge-v1").strip().lower()

        def is_better(new_att, old_att):
            if not old_att:
                return True
            n_score = new_att.get("score", 0)
            o_score = old_att.get("score", 0)
            if n_score > o_score:
                return True
            if n_score == o_score:
                n_time = new_att.get("timeTakenSeconds") or 999999
                o_time = old_att.get("timeTakenSeconds") or 999999
                return n_time < o_time
            return False

        if q_id in nmba_aliases or q_id == "nmba-awareness-challenge-v1":
            if is_better(a, participants_map[key]["nmba_best"]):
                participants_map[key]["nmba_best"] = a
        elif q_id in unity_aliases or q_id == "national-unity-day-quiz-v1":
            if is_better(a, participants_map[key]["unity_best"]):
                participants_map[key]["unity_best"] = a

    combined_list = []
    for key, data in participants_map.items():
        nmba_att = data["nmba_best"]
        unity_att = data["unity_best"]

        has_nmba = nmba_att is not None
        has_unity = unity_att is not None

        if eligibility == "both" and not (has_nmba and has_unity):
            continue
        if not (has_nmba or has_unity):
            continue

        completed_quizzes = []
        if has_nmba:
            completed_quizzes.append("nmba-awareness-challenge-v1")
        if has_unity:
            completed_quizzes.append("national-unity-day-quiz-v1")

        p = (nmba_att or unity_att).get("participant", {})
        nmba_score = nmba_att.get("score", 0) if has_nmba else 0
        nmba_max = nmba_att.get("maxScore", 100) if has_nmba else 100
        unity_score = unity_att.get("score", 0) if has_unity else 0
        unity_max = unity_att.get("maxScore", 100) if has_unity else 100

        total_score = nmba_score + unity_score
        total_max = (nmba_max if has_nmba else 0) + (unity_max if has_unity else 0)
        percentage = round((total_score / total_max) * 100) if total_max > 0 else 0

        nmba_time = nmba_att.get("timeTakenSeconds", 0) or 0 if has_nmba else 0
        unity_time = unity_att.get("timeTakenSeconds", 0) or 0 if has_unity else 0
        total_time = nmba_time + unity_time

        nmba_date = nmba_att.get("submittedAt", "") if has_nmba else ""
        unity_date = unity_att.get("submittedAt", "") if has_unity else ""
        latest_submitted = max(nmba_date, unity_date)

        # Assign badge
        if len(completed_quizzes) == 2:
            if percentage >= 85:
                badge = "NSS MJCET Dual Champion"
            else:
                badge = "National Solidarity Advocate"
        elif has_nmba:
            badge = "Campus Awareness Contributor"
        else:
            badge = "National Unity Contributor"

        combined_list.append({
            "attemptId": f"unified-{(nmba_att or unity_att).get('attemptId', key)}",
            "participantName": p.get("fullName", "Participant").strip(),
            "participantEmail": p.get("email", ""),
            "participantInstitution": p.get("institutionName", "MJCET").strip(),
            "rollNumber": p.get("rollNumber", ""),
            "phone": p.get("phone", ""),
            "department": p.get("department", ""),
            "yearOfStudy": p.get("yearOfStudy", ""),
            "score": total_score,
            "maxScore": total_max,
            "percentage": percentage,
            "passed": percentage >= 70,
            "timeTakenSeconds": total_time,
            "timeFormatted": format_duration_string(total_time),
            "badgeName": badge,
            "submittedAt": latest_submitted,
            "completedQuizzes": completed_quizzes,
            "quizzesAttemptedCount": len(completed_quizzes),
            "nashaMuktiScore": nmba_score if has_nmba else None,
            "unityDayScore": unity_score if has_unity else None,
            "nashaMuktiAttemptId": nmba_att.get("attemptId") if has_nmba else None,
            "unityDayAttemptId": unity_att.get("attemptId") if has_unity else None,
        })

    # Sort: -score, total_time, latest_submitted
    def sort_key(item):
        t = item["timeTakenSeconds"]
        t_val = 999999 if (t is None or t <= 0) else t
        return (-item["score"], t_val, item["submittedAt"])

    combined_list.sort(key=sort_key)

    ranked_items = []
    current_rank = 0
    prev_score = None
    prev_time = None

    for idx, doc in enumerate(combined_list):
        score = doc["score"]
        time_sec = doc["timeTakenSeconds"]

        is_tie = (prev_score is not None and score == prev_score and time_sec == prev_time)
        if is_tie:
            rank = current_rank
        else:
            current_rank = idx + 1
            rank = current_rank
            prev_score = score
            prev_time = time_sec

        doc["rank"] = rank
        ranked_items.append(doc)

    return ranked_items


async def compute_quiz_rankings(
    db: AsyncIOMotorDatabase,
    quiz_meta: Dict[str, Any],
) -> List[Dict[str, Any]]:
    """
    Computes authoritative rankings from genuine MongoDB quiz_attempts:
    1. Highest score first (-score).
    2. Shorter completion time ranks higher (ascending seconds).
    3. Handles exact ties transparently.
    """
    if quiz_meta["id"] == "unified-nss-mjcet-leaderboard":
        cfg = await get_or_create_leaderboard_record(db, quiz_meta)
        return await compute_unified_rankings(db, cfg)

    query_ids = quiz_meta.get("aliases", [quiz_meta["id"]])
    cursor = db["quiz_attempts"].find({"quizId": {"$in": query_ids}})
    attempts = []
    async for a in cursor:
        attempts.append(a)

    def rank_key(doc):
        score = doc.get("score", 0)
        time_sec = doc.get("timeTakenSeconds")
        if time_sec is None or time_sec <= 0:
            time_val = 999999
        else:
            time_val = time_sec
        sub = doc.get("submittedAt", "")
        return (-score, time_val, sub)

    attempts.sort(key=rank_key)

    ranked_items = []
    current_rank = 0
    prev_score = None
    prev_time = None

    for idx, doc in enumerate(attempts):
        score = doc.get("score", 0)
        time_sec = doc.get("timeTakenSeconds")

        is_tie = (prev_score is not None and score == prev_score and time_sec == prev_time)
        if is_tie:
            rank = current_rank
        else:
            current_rank = idx + 1
            rank = current_rank
            prev_score = score
            prev_time = time_sec

        part = doc.get("participant", {})
        ranked_items.append({
            "rank": rank,
            "attemptId": doc.get("attemptId", f"atmpt-{idx}"),
            "participantName": part.get("fullName", "Participant").strip(),
            "participantEmail": part.get("email", ""),
            "participantInstitution": part.get("institutionName", "MJCET").strip(),
            "rollNumber": part.get("rollNumber", ""),
            "phone": part.get("phone", ""),
            "department": part.get("department", ""),
            "yearOfStudy": part.get("yearOfStudy", ""),
            "score": score,
            "maxScore": doc.get("maxScore", 100),
            "percentage": doc.get("percentage", 0),
            "passed": doc.get("passed", False),
            "timeTakenSeconds": time_sec,
            "timeFormatted": format_duration_string(time_sec),
            "badgeName": doc.get("badgeName", "Participant"),
            "submittedAt": doc.get("submittedAt", ""),
            "completedQuizzes": [quiz_meta["id"]],
            "quizzesAttemptedCount": 1,
            "nashaMuktiScore": score if "nmba" in quiz_meta["id"] else None,
            "unityDayScore": score if "unity" in quiz_meta["id"] else None,
            "nashaMuktiAttemptId": doc.get("attemptId") if "nmba" in quiz_meta["id"] else None,
            "unityDayAttemptId": doc.get("attemptId") if "unity" in quiz_meta["id"] else None,
        })

    return ranked_items


@router.get("/quiz/leaderboards", response_model=List[LeaderboardSummaryItem])
async def list_public_quiz_leaderboards(
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Returns overview of public leaderboards for all active campaign quizzes and unified board."""
    results = []
    for q_key, meta in QUIZ_REGISTRY.items():
        cfg = await get_or_create_leaderboard_record(db, meta)
        if meta["id"] == "unified-nss-mjcet-leaderboard":
            # Count distinct participants across all attempts
            distinct_emails = await db["quiz_attempts"].distinct("participant.email")
            count = len([e for e in distinct_emails if e])
        else:
            count = await db["quiz_attempts"].count_documents({"quizId": {"$in": meta["aliases"]}})
        results.append(
            LeaderboardSummaryItem(
                quizId=meta["id"],
                quizTitle=meta["title"],
                quizSlug=meta["slug"],
                isPublished=cfg.get("isPublished", False),
                status="published" if cfg.get("isPublished", False) else "coming_soon",
                totalParticipants=count,
                announcementDate=cfg.get("announcementDate"),
            )
        )
    return results


@router.get("/quiz/leaderboard/{quiz_id}", response_model=LeaderboardPublicResponse)
async def get_public_quiz_leaderboard(
    quiz_id: str,
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Public leaderboard endpoint for a specific quiz or the unified leaderboard.
    STRICT SECURITY:
    If results are unpublished (isPublished == False), returns announcement state
    with 'entries = []'. Scores and rankings remain completely hidden.
    """
    meta = resolve_quiz_meta(quiz_id)
    cfg = await get_or_create_leaderboard_record(db, meta)
    is_published = cfg.get("isPublished", False)

    if meta["id"] == "unified-nss-mjcet-leaderboard":
        distinct_emails = await db["quiz_attempts"].distinct("participant.email")
        total_participants = len([e for e in distinct_emails if e])
    else:
        total_participants = await db["quiz_attempts"].count_documents({"quizId": {"$in": meta["aliases"]}})

    if not is_published:
        return LeaderboardPublicResponse(
            quizId=meta["id"],
            quizTitle=meta["title"],
            quizSlug=meta["slug"],
            isPublished=False,
            status="coming_soon",
            message="RESULTS TO BE ANNOUNCED SOON",
            supportingText=(
                "Our participants have taken on the challenge. "
                "Official results and rankings will be revealed soon by NSS MJCET. Stay tuned!"
            ),
            announcementDate=cfg.get("announcementDate"),
            customNote=cfg.get("customNote"),
            showPublicNames=cfg.get("showPublicNames", True),
            highlightTopThree=cfg.get("highlightTopThree", True),
            eligibilityRule=cfg.get("eligibilityRule", "any"),
            totalParticipants=total_participants,
            publishedAt=None,
            entries=[],
        )

    # Published state: compute official leaderboard
    if meta["id"] == "unified-nss-mjcet-leaderboard":
        ranked = await compute_unified_rankings(db, cfg)
    else:
        ranked = await compute_quiz_rankings(db, meta)

    show_names = cfg.get("showPublicNames", True)

    public_entries: List[LeaderboardEntryPublic] = []
    for item in ranked:
        name = item["participantName"] if show_names else mask_participant_name(item["participantName"])
        public_entries.append(
            LeaderboardEntryPublic(
                rank=item["rank"],
                attemptId=item["attemptId"],
                participantName=name,
                participantInstitution=item["participantInstitution"],
                score=item["score"],
                maxScore=item["maxScore"],
                percentage=item["percentage"],
                timeTakenSeconds=item["timeTakenSeconds"],
                timeFormatted=item["timeFormatted"],
                badgeName=item["badgeName"],
                submittedAt=item["submittedAt"],
                completedQuizzes=item.get("completedQuizzes"),
                nashaMuktiScore=item.get("nashaMuktiScore"),
                unityDayScore=item.get("unityDayScore"),
                quizzesAttemptedCount=item.get("quizzesAttemptedCount"),
            )
        )

    return LeaderboardPublicResponse(
        quizId=meta["id"],
        quizTitle=meta["title"],
        quizSlug=meta["slug"],
        isPublished=True,
        status="published",
        message="OFFICIAL RESULTS PUBLISHED",
        supportingText="Official certified rankings authenticated by NSS MJCET Examination Authority.",
        announcementDate=cfg.get("announcementDate"),
        customNote=cfg.get("customNote"),
        showPublicNames=show_names,
        highlightTopThree=cfg.get("highlightTopThree", True),
        eligibilityRule=cfg.get("eligibilityRule", "any"),
        totalParticipants=total_participants,
        publishedAt=cfg.get("publishedAt"),
        entries=public_entries,
    )


# =========================================================================
# Protected Admin Leaderboard Endpoints
# =========================================================================

@router.get("/admin/quiz/leaderboard/{quiz_id}", response_model=LeaderboardAdminResponse)
async def get_admin_quiz_leaderboard(
    quiz_id: str,
    current_admin: Dict[str, Any] = Depends(require_editor_or_above),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Returns full admin preview of calculated leaderboard for any quiz or the unified leaderboard.
    Available to Editors and Super Admins for verification.
    """
    meta = resolve_quiz_meta(quiz_id)
    cfg = await get_or_create_leaderboard_record(db, meta)

    if meta["id"] == "unified-nss-mjcet-leaderboard":
        ranked = await compute_unified_rankings(db, cfg)
    else:
        ranked = await compute_quiz_rankings(db, meta)

    total_participants = len(ranked)
    avg_score = round(sum(r["score"] for r in ranked) / total_participants, 1) if total_participants > 0 else 0
    highest_score = max((r["score"] for r in ranked), default=0)
    passed_count = sum(1 for r in ranked if r["passed"])

    admin_entries: List[LeaderboardEntryAdmin] = [
        LeaderboardEntryAdmin(**item) for item in ranked
    ]

    return LeaderboardAdminResponse(
        quizId=meta["id"],
        quizTitle=meta["title"],
        quizSlug=meta["slug"],
        isPublished=cfg.get("isPublished", False),
        status="published" if cfg.get("isPublished", False) else "draft",
        publishedAt=cfg.get("publishedAt"),
        publishedBy=cfg.get("publishedBy"),
        announcementDate=cfg.get("announcementDate"),
        showPublicNames=cfg.get("showPublicNames", True),
        highlightTopThree=cfg.get("highlightTopThree", True),
        eligibilityRule=cfg.get("eligibilityRule", "any"),
        customNote=cfg.get("customNote"),
        totalParticipants=total_participants,
        entries=admin_entries,
        stats={
            "totalParticipants": total_participants,
            "highestScore": highest_score,
            "averageScore": avg_score,
            "passedCount": passed_count,
            "passingRate": round((passed_count / total_participants) * 100, 1) if total_participants > 0 else 0,
        },
    )


@router.post("/admin/quiz/leaderboard/{quiz_id}/publish")
async def publish_quiz_leaderboard(
    quiz_id: str,
    current_admin: Dict[str, Any] = Depends(require_super_admin),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Authoritative publication of official quiz leaderboard.
    STRICT RBAC: Only Super Admins can publish results live to public users.
    """
    meta = resolve_quiz_meta(quiz_id)
    now_iso = datetime.now(timezone.utc).isoformat()

    await db["quiz_leaderboards"].update_one(
        {"quizId": meta["id"]},
        {
            "$set": {
                "quizTitle": meta["title"],
                "quizSlug": meta["slug"],
                "isPublished": True,
                "status": "published",
                "publishedAt": now_iso,
                "publishedBy": current_admin["email"],
                "updatedAt": now_iso,
            }
        },
        upsert=True,
    )

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="publish_quiz_leaderboard",
        module="quiz_leaderboard",
        record_id=meta["id"],
        metadata={"quizTitle": meta["title"]},
    )

    return {
        "success": True,
        "message": f"Official results for '{meta['title']}' published live successfully.",
        "publishedAt": now_iso,
        "publishedBy": current_admin["email"],
    }


@router.post("/admin/quiz/leaderboard/{quiz_id}/unpublish")
async def unpublish_quiz_leaderboard(
    quiz_id: str,
    current_admin: Dict[str, Any] = Depends(require_super_admin),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Retracts public leaderboard and returns to 'Results Coming Soon' state.
    STRICT RBAC: Only Super Admins can retract published results.
    """
    meta = resolve_quiz_meta(quiz_id)
    now_iso = datetime.now(timezone.utc).isoformat()

    await db["quiz_leaderboards"].update_one(
        {"quizId": meta["id"]},
        {
            "$set": {
                "isPublished": False,
                "status": "coming_soon",
                "updatedAt": now_iso,
            }
        },
        upsert=True,
    )

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="unpublish_quiz_leaderboard",
        module="quiz_leaderboard",
        record_id=meta["id"],
        metadata={"quizTitle": meta["title"]},
    )

    return {
        "success": True,
        "message": f"Official results for '{meta['title']}' have been retracted to 'Coming Soon'.",
    }


@router.put("/admin/quiz/leaderboard/{quiz_id}/settings")
async def update_quiz_leaderboard_settings(
    quiz_id: str,
    payload: LeaderboardSettingsUpdate,
    current_admin: Dict[str, Any] = Depends(require_super_admin),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """Updates presentation settings (announcement date, public names, top 3 highlight)."""
    meta = resolve_quiz_meta(quiz_id)
    now_iso = datetime.now(timezone.utc).isoformat()

    update_fields: Dict[str, Any] = {"updatedAt": now_iso}
    if payload.announcementDate is not None:
        update_fields["announcementDate"] = payload.announcementDate.strip()
    if payload.showPublicNames is not None:
        update_fields["showPublicNames"] = payload.showPublicNames
    if payload.highlightTopThree is not None:
        update_fields["highlightTopThree"] = payload.highlightTopThree
    if payload.customNote is not None:
        update_fields["customNote"] = payload.customNote.strip()
    if payload.eligibilityRule is not None:
        update_fields["eligibilityRule"] = payload.eligibilityRule

    await db["quiz_leaderboards"].update_one(
        {"quizId": meta["id"]},
        {"$set": update_fields},
        upsert=True,
    )

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="update_quiz_leaderboard_settings",
        module="quiz_leaderboard",
        record_id=meta["id"],
        metadata=update_fields,
    )

    return {"success": True, "message": "Leaderboard settings updated successfully."}


@router.put("/admin/quiz/attempts/{attempt_id}/override")
async def override_quiz_attempt(
    attempt_id: str,
    payload: AttemptOverrideRequest,
    current_admin: Dict[str, Any] = Depends(require_super_admin),
    db: AsyncIOMotorDatabase = Depends(get_database),
):
    """
    Authoritative administrative override for participant score, duration, or identity correction.
    STRICT RBAC: Only Super Admins can alter recorded attempt results.
    """
    target = await db["quiz_attempts"].find_one({"attemptId": attempt_id})
    if not target:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Attempt '{attempt_id}' was not found in records.",
        )

    now_iso = datetime.now(timezone.utc).isoformat()
    update_doc: Dict[str, Any] = {"updatedAt": now_iso, "lastModifiedBy": current_admin["email"]}

    if payload.score is not None:
        update_doc["score"] = payload.score
        max_s = target.get("maxScore", 100)
        pct = round((payload.score / max_s) * 100) if max_s > 0 else 0
        update_doc["percentage"] = pct
        update_doc["passed"] = pct >= 70
    if payload.timeTakenSeconds is not None:
        update_doc["timeTakenSeconds"] = payload.timeTakenSeconds
    if payload.participantName is not None and payload.participantName.strip():
        update_doc["participant.fullName"] = payload.participantName.strip()
    if payload.notes:
        update_doc["adminNotes"] = payload.notes.strip()

    await db["quiz_attempts"].update_one({"attemptId": attempt_id}, {"$set": update_doc})

    await record_audit_log(
        db,
        actor_email=current_admin["email"],
        action="override_quiz_attempt",
        module="quiz_attempts",
        record_id=attempt_id,
        metadata={"changes": list(update_doc.keys()), "actor": current_admin["email"]},
    )

    return {"success": True, "message": f"Attempt '{attempt_id}' updated successfully."}
